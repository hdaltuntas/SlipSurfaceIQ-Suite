"""
Calculation report for SlipSurface IQ Rock.

The report is assembled as HTML (inputs, the Hoek–Brown constants and
strengths, the equivalent Mohr–Coulomb strength of every application, the
deformation modulus by every estimate, the instantaneous and residual
strength, the figures, the warnings, the method notes and, if one was run,
the study) and exported as

  * PDF   — through reportlab (see `slipsurface.common.pdf`)
  * HTML  — a single self-contained file (figures embedded as base64)
  * DOCX  — optional, needs `python-docx`

There is one assembly, `build_html()`, so all three formats say the same thing.
"""

from __future__ import annotations

import base64
import datetime
import html
import io
import math
from typing import Any, Dict, List, Optional

from matplotlib.figure import Figure

from . import study_plots, tables
from .config import APP_NAME, APP_VERSION, MODULUS_METHODS
from .i18n import TRANSLATIONS, warning_text
from .plotting import Plotter, available_figures
from .study import CHARACTERISTIC, OUTPUTS
from .summary import gsi_source_text

TEXTS = {
    "en": {
        "title": f"{APP_NAME} — Calculation Report",
        "date": "Date", "analyst": "Analyst", "software": "Software",
        "sec_inputs": "1. Input data", "sec_intact": "1.1 Intact rock",
        "sec_mass": "1.2 Rock mass", "sec_app": "1.3 Application and deformability",
        "sec_lab": "1.4 Laboratory tests",
        "sec_hb": "2. Generalised Hoek–Brown criterion", "sec_const": "2.1 Constants",
        "sec_strength": "2.2 Rock mass strength",
        "sec_mc": "3. Equivalent Mohr–Coulomb strength",
        "sec_mc_used": "3.1 For the application in use",
        "sec_mc_all": "3.2 For every application",
        "sec_modulus": "4. Deformation modulus",
        "sec_query": "5. Instantaneous strength",
        "sec_residual": "6. Residual strength (Cai et al. 2007)",
        "sec_figs": "7. Figures", "sec_warn": "8. Warnings", "sec_notes": "9. Method notes",
        "parameter": "Parameter", "value": "Value", "unit": "Unit",
        "rock_type": "Rock type", "source": "Source of σci and mi",
        "grade": "Field strength grade", "sigma_ci": "Uniaxial strength of intact rock σci",
        "mi": "Hoek–Brown constant mi", "Ei": "Intact modulus Ei", "MR": "Modulus ratio MR",
        "Ei_source": "Intact modulus from",
        "gsi": "Geological Strength Index GSI", "gsi_how": "GSI from",
        "D": "Disturbance factor D", "d_case": "Excavation",
        "application": "Application", "gamma": "Unit weight of the rock mass γ",
        "depth": "Depth of the tunnel / height of the slope H", "gammaH": "γH",
        "s3max": "Upper limit of confinement σ3max", "modulus_method": "Modulus estimate",
        "query_s3": "σ3 of the instantaneous strength",
        "test": "Test", "fit_sci": "Fitted σci", "fit_mi": "Fitted mi", "r2": "r²",
        "fit_range": "σ3 range of the tests",
        "mb": "mb", "s": "s", "a": "a", "mb_mi": "mb/mi",
        "sigma_t": "Tensile strength σt", "sigma_c": "Uniaxial compressive strength σc",
        "sigma_cm": "Global rock mass strength σcm", "sigma_cm_ci": "σcm/σci",
        "c": "Cohesion c'", "phi": "Friction angle φ'",
        "mc_ucs": "Uniaxial strength of the fitted line", "mc_t": "Tensile strength of the fitted line",
        "k": "Passive ratio (1 + sin φ')/(1 − sin φ')",
        "method": "Method", "Erm": "Erm", "Erm_Ei": "Erm/Ei",
        "sigma1": "σ1", "sigma_n": "σn", "tau": "τ", "phi_i": "Instantaneous friction angle φi",
        "c_i": "Instantaneous cohesion ci",
        "peak": "Peak", "residual": "Residual",
        "notes": [
            "Generalised Hoek–Brown criterion, 2002 edition (Hoek, Carranza-Torres & Corkum "
            "2002): σ1 = σ3 + σci(mb·σ3/σci + s)^a with mb = mi·e^((GSI−100)/(28−14D)), "
            "s = e^((GSI−100)/(9−3D)), a = ½ + (e^(−GSI/15) − e^(−20/3))/6.",
            "The criterion describes a rock mass that behaves isotropically: jointed closely "
            "enough, relative to the size of the structure, that no single discontinuity "
            "governs. Where one does, analyse that discontinuity instead.",
            "The equivalent Mohr–Coulomb strength is the line fitted to the Hoek–Brown curve "
            "over σt < σ3 < σ3max (Hoek et al. 2002). σ3max is σci/4 for a general "
            "application; 0.47·σcm·(σcm/γH)^−0.94 for a tunnel and 0.72·σcm·(σcm/γH)^−0.91 "
            "for a slope, H being the depth of the tunnel or the height of the slope.",
            "Deformation modulus: Hoek & Diederichs (2006), generalised "
            "Erm = Ei(0.02 + (1 − D/2)/(1 + e^((60 + 15D − GSI)/11))) and simplified "
            "Erm = 100 000(1 − D/2)/(1 + e^((75 + 25D − GSI)/11)) MPa; Hoek et al. (2002) "
            "Erm = (1 − D/2)√(σci/100)·10^((GSI − 10)/40) GPa.",
            "Instantaneous strength after Balmer (1952): the tangent to the envelope in the "
            "Mohr plane at the chosen σ3.",
            "Residual strength after Cai et al. (2007): GSIr = GSI·e^(−0.0134·GSI), with the "
            "intact rock's σci and mi.",
            "Typical values of mi and MR from Marinos & Hoek (2000) and Hoek & Diederichs "
            "(2006); field strength grades after ISRM, as given by Hoek (2007).",
            "D applies to the rock damaged by the excavation, which in a tunnel is a zone a "
            "few metres thick; the rock beyond it is undisturbed (D = 0).",
        ],
        "note_gsi": {
            "chart": "GSI read from the classes of the GSI chart, placed on the quantified "
                     "chart of Hoek, Carter & Diederichs (2013).",
            "jcond": "GSI = 1.5·JCond89 + RQD/2 (Hoek, Carter & Diederichs 2013).",
            "rmr": "GSI = RMR89 − 5, RMR89 with a dry groundwater rating and no orientation "
                   "adjustment (Hoek, Kaiser & Bawden 1995); not for RMR89 < 23.",
            "q": "GSI = 9·ln Q' + 44, Q' = RQD/Jn · Jr/Ja (Hoek et al. 1995).",
            "cai": "GSI from the block volume Vb and the joint condition factor Jc "
                   "(Cai et al. 2004).",
        },
        "note_lab": "σci and mi fitted by least squares to (σ1 − σ3)² = mi·σci·σ3 + σci² "
                    "(Hoek & Brown 1980); tests over 0 ≤ σ3 ≤ σci/2 are recommended.",
        "yes": "yes", "no": "no",
        "sec_study": "10. Parametric / probabilistic study", "sec_study_vars": "10.1 Variables",
        "sec_study_stats": "10.2 Statistics of the outputs",
        "sec_study_char": "10.3 Characteristic values (5 % fractile)",
        "sec_study_sens": "10.4 Sensitivities (Spearman ρ)", "sec_study_figs": "10.5 Figures",
        "study_method": "Sampling", "study_n": "Samples", "study_ok": "Analysed",
        "variable": "Variable", "mode": "Mode", "range": "range", "dist": "distribution",
        "min": "Min", "max": "Max", "mean": "Mean", "cov": "CoV", "std": "Std. dev.",
        "output": "Output", "p5": "P5", "ratio": "P5/mean",
    },
    "tr": {
        "title": f"{APP_NAME} — Hesap Raporu",
        "date": "Tarih", "analyst": "Hazırlayan", "software": "Yazılım",
        "sec_inputs": "1. Girdi verileri", "sec_intact": "1.1 Sağlam kaya",
        "sec_mass": "1.2 Kaya kütlesi", "sec_app": "1.3 Uygulama ve deformabilite",
        "sec_lab": "1.4 Laboratuvar deneyleri",
        "sec_hb": "2. Genelleştirilmiş Hoek–Brown ölçütü", "sec_const": "2.1 Sabitler",
        "sec_strength": "2.2 Kaya kütlesi dayanımı",
        "sec_mc": "3. Eşdeğer Mohr–Coulomb dayanımı",
        "sec_mc_used": "3.1 Kullanılan uygulama için",
        "sec_mc_all": "3.2 Tüm uygulamalar için",
        "sec_modulus": "4. Deformasyon modülü",
        "sec_query": "5. Anlık dayanım",
        "sec_residual": "6. Kalıcı (rezidüel) dayanım (Cai vd. 2007)",
        "sec_figs": "7. Şekiller", "sec_warn": "8. Uyarılar", "sec_notes": "9. Yöntem notları",
        "parameter": "Parametre", "value": "Değer", "unit": "Birim",
        "rock_type": "Kaya türü", "source": "σci ve mi kaynağı",
        "grade": "Arazi dayanım sınıfı", "sigma_ci": "Sağlam kayanın tek eksenli dayanımı σci",
        "mi": "Hoek–Brown sabiti mi", "Ei": "Sağlam kaya modülü Ei", "MR": "Modül oranı MR",
        "Ei_source": "Sağlam kaya modülü",
        "gsi": "Jeolojik Dayanım İndeksi GSI", "gsi_how": "GSI kaynağı",
        "D": "Örselenme faktörü D", "d_case": "Kazı",
        "application": "Uygulama", "gamma": "Kaya kütlesi birim hacim ağırlığı γ",
        "depth": "Tünel derinliği / şev yüksekliği H", "gammaH": "γH",
        "s3max": "Çevre basıncının üst sınırı σ3max", "modulus_method": "Modül tahmini",
        "query_s3": "Anlık dayanımın σ3 değeri",
        "test": "Deney", "fit_sci": "Uydurulan σci", "fit_mi": "Uydurulan mi", "r2": "r²",
        "fit_range": "Deneylerin σ3 aralığı",
        "mb": "mb", "s": "s", "a": "a", "mb_mi": "mb/mi",
        "sigma_t": "Çekme dayanımı σt", "sigma_c": "Tek eksenli basınç dayanımı σc",
        "sigma_cm": "Global kaya kütlesi dayanımı σcm", "sigma_cm_ci": "σcm/σci",
        "c": "Kohezyon c'", "phi": "İçsel sürtünme açısı φ'",
        "mc_ucs": "Uydurulan doğrunun tek eksenli dayanımı",
        "mc_t": "Uydurulan doğrunun çekme dayanımı",
        "k": "Pasif oran (1 + sin φ')/(1 − sin φ')",
        "method": "Yöntem", "Erm": "Erm", "Erm_Ei": "Erm/Ei",
        "sigma1": "σ1", "sigma_n": "σn", "tau": "τ", "phi_i": "Anlık sürtünme açısı φi",
        "c_i": "Anlık kohezyon ci",
        "peak": "Pik", "residual": "Kalıcı",
        "notes": [
            "Genelleştirilmiş Hoek–Brown ölçütü, 2002 sürümü (Hoek, Carranza-Torres & Corkum "
            "2002): σ1 = σ3 + σci(mb·σ3/σci + s)^a; mb = mi·e^((GSI−100)/(28−14D)), "
            "s = e^((GSI−100)/(9−3D)), a = ½ + (e^(−GSI/15) − e^(−20/3))/6.",
            "Ölçüt izotrop davranan bir kaya kütlesini tanımlar: yapının boyutuna göre "
            "yeterince sık eklemli, tek bir süreksizliğin davranışı yönetmediği kütle. "
            "Yöneten bir süreksizlik varsa o süreksizlik ayrıca incelenmelidir.",
            "Eşdeğer Mohr–Coulomb dayanımı, Hoek–Brown eğrisine σt < σ3 < σ3max aralığında "
            "uydurulan doğrudur (Hoek vd. 2002). σ3max genel uygulamada σci/4; tünelde "
            "0.47·σcm·(σcm/γH)^−0.94, şevde 0.72·σcm·(σcm/γH)^−0.91'dir; H tünel derinliği "
            "ya da şev yüksekliğidir.",
            "Deformasyon modülü: Hoek & Diederichs (2006) genelleştirilmiş "
            "Erm = Ei(0.02 + (1 − D/2)/(1 + e^((60 + 15D − GSI)/11))) ve basitleştirilmiş "
            "Erm = 100 000(1 − D/2)/(1 + e^((75 + 25D − GSI)/11)) MPa; Hoek vd. (2002) "
            "Erm = (1 − D/2)√(σci/100)·10^((GSI − 10)/40) GPa.",
            "Anlık dayanım Balmer (1952) yöntemiyle: seçilen σ3'te Mohr düzlemindeki "
            "zarfa teğet.",
            "Kalıcı dayanım Cai vd. (2007): GSIr = GSI·e^(−0.0134·GSI), sağlam kayanın σci "
            "ve mi değerleriyle.",
            "Tipik mi ve MR değerleri Marinos & Hoek (2000) ve Hoek & Diederichs (2006); "
            "arazi dayanım sınıfları ISRM'ye göre, Hoek (2007) aktarımıyla.",
            "D yalnızca kazının hasar verdiği kaya için geçerlidir; tünelde bu birkaç metre "
            "kalınlıkta bir bölgedir, ötesindeki kaya örselenmemiştir (D = 0).",
        ],
        "note_gsi": {
            "chart": "GSI, GSI abağının sınıflarından okunmuş ve Hoek, Carter & Diederichs "
                     "(2013) sayısallaştırılmış abağına yerleştirilmiştir.",
            "jcond": "GSI = 1.5·JCond89 + RQD/2 (Hoek, Carter & Diederichs 2013).",
            "rmr": "GSI = RMR89 − 5; RMR89 kuru yeraltı suyu puanıyla ve yönelim düzeltmesi "
                   "olmadan (Hoek, Kaiser & Bawden 1995); RMR89 < 23 için kullanılmaz.",
            "q": "GSI = 9·ln Q' + 44, Q' = RQD/Jn · Jr/Ja (Hoek vd. 1995).",
            "cai": "GSI, blok hacmi Vb ve süreksizlik durum faktörü Jc'den (Cai vd. 2004).",
        },
        "note_lab": "σci ve mi, (σ1 − σ3)² = mi·σci·σ3 + σci² doğrusuna en küçük kareler "
                    "yöntemiyle uydurulmuştur (Hoek & Brown 1980); deneylerin "
                    "0 ≤ σ3 ≤ σci/2 aralığında olması önerilir.",
        "yes": "evet", "no": "hayır",
        "sec_study": "10. Parametrik / olasılıksal çalışma", "sec_study_vars": "10.1 Değişkenler",
        "sec_study_stats": "10.2 Çıktıların istatistikleri",
        "sec_study_char": "10.3 Karakteristik değerler (%5 alt değer)",
        "sec_study_sens": "10.4 Duyarlılıklar (Spearman ρ)", "sec_study_figs": "10.5 Şekiller",
        "study_method": "Örnekleme", "study_n": "Örnek sayısı", "study_ok": "Çözülen",
        "variable": "Değişken", "mode": "Tür", "range": "aralık", "dist": "dağılım",
        "min": "En az", "max": "En çok", "mean": "Ortalama", "cov": "DK", "std": "Std. sapma",
        "output": "Çıktı", "p5": "P5", "ratio": "P5/ortalama",
    },
}

_CSS = """
body { font-family: 'Segoe UI', 'Helvetica Neue', Arial, 'DejaVu Sans', sans-serif;
       font-size: 10pt; color: #141413; max-width: 900px; margin: 24px auto; }
h1, h2, h3, .lead { font-family: 'Tiempos Text', 'Source Serif 4', 'Palatino Linotype',
             Georgia, 'DejaVu Serif', serif; font-weight: 500; }
h1 { font-size: 20pt; color: #141413; margin-bottom: 2px; }
h2 { font-size: 14pt; color: #C6613F; border-bottom: 1px solid #E3E0D5; padding-bottom: 2px;
     margin-top: 20px; }
h3 { font-size: 11.5pt; color: #141413; margin-top: 12px; }
table { border-collapse: collapse; margin: 4px 0 8px 0; }
th { background: #F0EEE6; text-align: left; padding: 3px 6px; border: 1px solid #E3E0D5;
     font-size: 9pt; }
td { padding: 3px 6px; border: 1px solid #E3E0D5; font-size: 9pt; }
.meta { color: #73726C; } .note { color: #73726C; font-size: 9pt; }
"""


def _esc(x) -> str:
    return html.escape(str(x))


def _f(x, nd=3) -> str:
    """A number with `nd` significant figures, thousands separated."""
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return "—"
    if x == 0:
        return "0"
    digits = max(nd - 1 - int(math.floor(math.log10(abs(x)))), 0)
    return f"{x:,.{digits}f}"


def _table(headers: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None) -> str:
    out = ["<table width='100%'>"]
    if widths:
        cells = "".join(f"<th width='{w}%'>{_esc(h)}</th>" for h, w in zip(headers, widths))
    else:
        cells = "".join(f"<th>{_esc(h)}</th>" for h in headers)
    out.append("<tr>" + cells + "</tr>")
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def _kv_table(T, rows: List[List[Any]]) -> str:
    return _table([T["parameter"], T["value"], T["unit"]], rows, [54, 32, 14])


# ----------------------------------------------------------------------
def _png(fig: Figure, dpi: int) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    return buf.getvalue()


def render_figures(analysis, lang: str, keys=None, dpi: int = 130) -> Dict[str, bytes]:
    """The report figures as PNG bytes (off-screen, white, untitled)."""
    plotter = Plotter(analysis, lang, "paper", titles=False)
    out = {}
    for key in keys or available_figures(analysis):
        size = (11, 5.2) if key in ("envelopes", "disturbance", "labfit") else (10, 6.5)
        fig = Figure(figsize=size, dpi=dpi)
        plotter.draw(key, fig)
        out[key] = _png(fig, dpi)
    return out


def render_study_figures(study, lang: str, dpi: int = 110) -> Dict[str, bytes]:
    """Study figures keyed 'study_<view>' (only the views that apply)."""
    if study is None or not study.rows:
        return {}
    L = TRANSLATIONS[lang]
    views = ["oat", "hist"] if study.method == "oat" else ["hist", "scatter", "tornado"]
    draw = {"oat": study_plots.plot_oat, "hist": study_plots.plot_hist,
            "scatter": study_plots.plot_scatter, "tornado": study_plots.plot_tornado}
    out = {}
    for view in views:
        fig = Figure(figsize=(10, 7), dpi=dpi)
        draw[view](fig, study, L, "phi", theme="paper")
        out[f"study_{view}"] = _png(fig, dpi)
    return out


def _study_section(T, L, study, figures, img_src) -> List[str]:
    parts = [f"<h2 style='page-break-before:always'>{T['sec_study']}</h2>"]
    s = study.summary
    parts.append(_kv_table(T, [
        [T["study_method"], L[f"sampling_{study.method}"], ""],
        [T["study_n"], str(s.get("n_total", 0)), ""],
        [T["study_ok"], str(s.get("n_ok", 0)), ""],
    ]))
    parts.append(f"<h3>{T['sec_study_vars']}</h3>")
    rows = []
    for v in study.variables:
        if v.mode == "range":
            rows.append([_esc(v.label), T["range"], _f(v.vmin, 4), _f(v.vmax, 4), "—", "—",
                         "—", v.n_points])
        else:
            rows.append([_esc(v.label), T["dist"], "—", "—", L[f"dist_{v.dist}"],
                         _f(v.mean, 4), f"{v.cov:.3f}", "—"])
    parts.append(_table([T["variable"], T["mode"], T["min"], T["max"], T["dist"], T["mean"],
                         T["cov"], "n"], rows, [28, 12, 10, 10, 14, 10, 8, 8]))
    stats = s.get("stats", {})
    parts.append(f"<h3>{T['sec_study_stats']}</h3>")
    rows = [[_esc(L.get(f"out_{k}", k)), st["n"], _f(st["mean"], 4), _f(st["std"], 3),
             _f(st["cov"], 3), _f(st["p5"], 4), _f(st["p50"], 4), _f(st["p95"], 4)]
            for k, _ in OUTPUTS if k in stats for st in [stats[k]]]
    parts.append(_table([T["output"], "n", T["mean"], T["std"], T["cov"], "P5", "P50", "P95"],
                        rows, [28, 7, 11, 11, 8, 11, 12, 12]))
    chars = [k for k in CHARACTERISTIC if k in stats]
    if chars and study.method != "oat":
        parts.append(f"<h3>{T['sec_study_char']}</h3>")
        rows = [[_esc(L.get(f"out_{k}", k)), _f(stats[k]["mean"], 4), _f(stats[k]["p5"], 4),
                 f"{stats[k]['p5'] / stats[k]['mean']:.3f}" if stats[k]["mean"] else "—"]
                for k in chars]
        parts.append(_table([T["output"], T["mean"], T["p5"], T["ratio"]], rows,
                            [40, 20, 20, 20]))
    if s.get("spearman"):
        parts.append(f"<h3>{T['sec_study_sens']}</h3>")
        outs = [k for k in ("c", "phi", "Erm", "sigma_cm") if k in s["spearman"]]
        rows = [[_esc(v.label)] + [f"{s['spearman'][k].get(v.path, float('nan')):+.2f}"
                                   for k in outs] for v in study.variables]
        parts.append(_table([T["variable"]] + [L.get(f"out_{k}", k) for k in outs], rows))
    parts.append(f"<h3>{T['sec_study_figs']}</h3>")
    for key in figures:
        if key.startswith("study_"):
            parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")
    return parts


def build_html(analysis, lang: str, figures: Dict[str, bytes], img_src=None,
               study=None) -> str:
    """
    Builds the report HTML. img_src(key) -> value for the <img src> attribute;
    the default embeds base64 data URIs (the self-contained HTML file). The PDF
    and DOCX writers pass a `fig://<key>` mapper and resolve the keys against
    the figure dictionary themselves.
    """
    T = TEXTS.get(lang, TEXTS["en"])
    L = TRANSLATIONS.get(lang, TRANSLATIONS["en"])
    a, res = analysis, analysis.results
    cfg = a.config
    info = cfg.get("project_info", {})
    if img_src is None:
        def img_src(key):
            return "data:image/png;base64," + base64.b64encode(figures[key]).decode("ascii")

    intact, h, st, mc, m, q, r = (res["intact"], res["hb"], res["strength"], res["mc"],
                                  res["modulus"], res["query"], res["residual"])

    parts = [f"<html><head><meta charset='utf-8'><style>{_CSS}</style></head><body>"]
    parts.append(f"<h1>{_esc(info.get('title') or T['title'])}</h1>")
    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    meta = f"{T['date']}: {stamp}"
    if info.get("analyst"):
        meta += f" &nbsp;·&nbsp; {T['analyst']}: {_esc(info['analyst'])}"
    meta += f" &nbsp;·&nbsp; {T['software']}: {APP_NAME} v{APP_VERSION}"
    parts.append(f"<p class='meta'>{meta}</p>")

    # ---------------------------------------------------------------- inputs
    parts.append(f"<h2>{T['sec_inputs']}</h2>")
    parts.append(f"<h3>{T['sec_intact']}</h3>")
    rows = []
    if tables.rock_type(intact["rock_type"]) is not None:
        rows.append([T["rock_type"],
                     _esc(tables.name(tables.ROCK_TYPES, intact["rock_type"], lang)), ""])
    rows.append([T["source"], L[f"source_{intact['source']}"], ""])
    grade = cfg["intact"].get("strength_grade")
    if intact["source"] == "manual" and tables.strength_grade(grade) is not None:
        rows.append([T["grade"], _esc(tables.name(tables.STRENGTH_GRADES, grade, lang)), ""])
    rows += [[T["sigma_ci"], _f(intact["sigma_ci"], 4), "MPa"],
             [T["mi"], _f(intact["mi"], 4), ""],
             [T["Ei_source"], L[f"Ei_source_{cfg['modulus']['Ei_source']}"], ""],
             [T["MR"], _f(intact["MR"], 4), ""],
             [T["Ei"], _f(intact["Ei"], 4), "MPa"]]
    parts.append(_kv_table(T, rows))

    parts.append(f"<h3>{T['sec_mass']}</h3>")
    rows = [[T["gsi"], f"{res['GSI']:.1f}", ""],
            [T["gsi_how"], _esc(gsi_source_text(L, res, lang)), ""],
            [T["D"], f"{res['D']:.2f}", ""]]
    case = cfg["rock_mass"].get("disturbance_case")
    if tables.disturbance(case) is not None:
        rows.append([T["d_case"], _esc(tables.name(tables.DISTURBANCE, case, lang)), ""])
    parts.append(_kv_table(T, rows))

    parts.append(f"<h3>{T['sec_app']}</h3>")
    rows = [[T["application"], L[f"application_{res['application']}"], ""]]
    if res["application"] in ("tunnel", "slope"):
        rows += [[T["gamma"], _f(cfg["application"]["unit_weight"], 3), "kN/m³"],
                 [T["depth"], _f(cfg["application"]["depth"], 4), "m"],
                 [T["gammaH"], _f(res["gamma_H"], 4), "MPa"]]
    rows += [[T["s3max"], _f(res["sigma3_max"], 4), "MPa"],
             [T["modulus_method"], L[f"modulus_{m['method']}"], ""],
             [T["query_s3"], _f(q["sigma3"], 4), "MPa"]]
    parts.append(_kv_table(T, rows))

    fit = intact.get("fit")
    if fit:
        parts.append(f"<h3>{T['sec_lab']}</h3>")
        rows = [[_esc(t["name"] or f"#{i + 1}"), _f(t["sigma3"], 4), _f(t["sigma1"], 4)]
                for i, t in enumerate(fit["tests"])]
        parts.append(_table([T["test"], "σ3 [MPa]", "σ1 [MPa]"], rows, [40, 30, 30]))
        parts.append(_kv_table(T, [
            [T["fit_sci"], _f(fit["sigma_ci"], 4), "MPa"],
            [T["fit_mi"], _f(fit["mi"], 4), ""],
            [T["r2"], f"{fit['r2']:.4f}", ""],
            [T["fit_range"], f"{_f(fit['sigma3_min'])} … {_f(fit['sigma3_max'])}", "MPa"],
        ]))

    # ---------------------------------------------------------- Hoek–Brown
    parts.append(f"<h2>{T['sec_hb']}</h2>")
    parts.append(f"<h3>{T['sec_const']}</h3>")
    parts.append(_kv_table(T, [
        [T["mb"], _f(h["mb"], 5), ""], [T["s"], f"{h['s']:.5g}", ""],
        [T["a"], f"{h['a']:.5f}", ""], [T["mb_mi"], f"{h['mb_mi']:.4f}", ""],
    ]))
    parts.append(f"<h3>{T['sec_strength']}</h3>")
    parts.append(_kv_table(T, [
        [T["sigma_t"], _f(st["sigma_t"], 4), "MPa"],
        [T["sigma_c"], _f(st["sigma_c"], 4), "MPa"],
        [T["sigma_cm"], _f(st["sigma_cm"], 4), "MPa"],
        [T["sigma_cm_ci"], f"{st['sigma_cm'] / intact['sigma_ci']:.4f}", ""],
    ]))

    # ---------------------------------------------------------- Mohr–Coulomb
    parts.append(f"<h2>{T['sec_mc']}</h2>")
    parts.append(f"<h3>{T['sec_mc_used']} — {L['application_' + res['application']]}</h3>")
    parts.append(_kv_table(T, [
        [T["s3max"], _f(res["sigma3_max"], 4), "MPa"],
        [T["c"], _f(mc["c"], 4), "MPa"],
        [T["phi"], f"{mc['phi']:.2f}", "°"],
        [T["k"], f"{mc['k']:.4f}", ""],
        [T["mc_ucs"], _f(mc["sigma_cm"], 4), "MPa"],
        [T["mc_t"], _f(mc["sigma_t"], 4), "MPa"],
    ]))
    parts.append(f"<h3>{T['sec_mc_all']}</h3>")
    rows = []
    for name, entry in res["applications"].items():
        mark = " *" if name == res["application"] else ""
        rows.append([_esc(L[f"application_{name}"] + mark), _f(entry["sigma3_max"], 4),
                     _f(entry["c"], 4), f"{entry['phi']:.2f}"])
    parts.append(_table([T["application"], "σ3max [MPa]", "c' [MPa]", "φ' [°]"], rows,
                        [40, 20, 20, 20]))

    # ---------------------------------------------------------- modulus
    parts.append(f"<h2>{T['sec_modulus']}</h2>")
    rows = []
    for name in MODULUS_METHODS:
        mark = " *" if name == m["method"] else ""
        rows.append([_esc(L[f"modulus_{name}"] + mark), _f(m["all"][name], 4),
                     f"{m['all'][name] / m['Ei']:.4f}"])
    parts.append(_table([T["method"], "Erm [MPa]", T["Erm_Ei"]], rows, [50, 25, 25]))

    # ---------------------------------------------------------- instantaneous
    parts.append(f"<h2>{T['sec_query']}</h2>")
    parts.append(_kv_table(T, [
        ["σ3", _f(q["sigma3"], 4), "MPa"], [T["sigma1"], _f(q["sigma1"], 4), "MPa"],
        [T["sigma_n"], _f(q["sigma_n"], 4), "MPa"], [T["tau"], _f(q["tau"], 4), "MPa"],
        [T["phi_i"], f"{q['phi']:.2f}", "°"], [T["c_i"], _f(q["c"], 4), "MPa"],
    ]))

    # ---------------------------------------------------------- residual
    parts.append(f"<h2>{T['sec_residual']}</h2>")
    rows = [["GSI", f"{res['GSI']:.1f}", f"{r['GSI']:.1f}"],
            ["mb", _f(h["mb"], 4), _f(r["mb"], 4)],
            ["s", f"{h['s']:.4g}", f"{r['s']:.4g}"],
            ["a", f"{h['a']:.4f}", f"{r['a']:.4f}"],
            ["c' [MPa]", _f(mc["c"], 4), _f(r["c"], 4)],
            ["φ' [°]", f"{mc['phi']:.2f}", f"{r['phi']:.2f}"],
            ["σt [MPa]", _f(st["sigma_t"], 4), _f(r["sigma_t"], 4)],
            ["σc [MPa]", _f(st["sigma_c"], 4), _f(r["sigma_c"], 4)]]
    parts.append(_table([T["parameter"], T["peak"], T["residual"]], rows, [40, 30, 30]))

    # ---------------------------------------------------------- figures
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_figs']}</h2>")
    for key in available_figures(analysis):
        if key not in figures:
            continue
        parts.append(f"<p class='lead'>{_esc(L.get('fig_' + key, key))}</p>")
        parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")

    if res["warnings"]:
        parts.append(f"<h2>{T['sec_warn']}</h2><ul>")
        parts.extend(f"<li>{_esc(warning_text(lang, w))}</li>" for w in res["warnings"])
        parts.append("</ul>")

    parts.append(f"<h2>{T['sec_notes']}</h2><ul class='note'>")
    notes = list(T["notes"])
    if res["gsi"]["source"] in T["note_gsi"]:
        notes.append(T["note_gsi"][res["gsi"]["source"]])
    if fit:
        notes.append(T["note_lab"])
    parts.extend(f"<li>{_esc(n)}</li>" for n in notes)
    parts.append("</ul>")

    if study is not None and study.rows:
        parts.extend(_study_section(T, L, study, figures, img_src))
    parts.append("</body></html>")
    return "\n".join(parts)


# ----------------------------------------------------------------------
# exporters
# ----------------------------------------------------------------------
def _all_figures(analysis, lang, study):
    figures = render_figures(analysis, lang)
    figures.update(render_study_figures(study, lang))
    return figures


def export_html(path: str, analysis, lang: str, study=None) -> None:
    figures = _all_figures(analysis, lang, study)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(build_html(analysis, lang, figures, study=study))


def export_pdf(path: str, analysis, lang: str, study=None) -> None:
    """PDF through reportlab; the figures travel as `fig://<key>` references."""
    from ..common import pdf as pdf_writer

    T = TEXTS.get(lang, TEXTS["en"])
    info = analysis.config.get("project_info", {})
    figures = _all_figures(analysis, lang, study)
    html_text = build_html(analysis, lang, figures, img_src=lambda k: f"fig://{k}",
                           study=study)
    pdf_writer.html_to_pdf(path, html_text, figures,
                           title=info.get("title", T["title"]),
                           author=info.get("analyst", ""),
                           footer=f"{info.get('title', '')} — {APP_NAME} v{APP_VERSION}")


def export_docx(path: str, analysis, lang: str, study=None) -> None:
    """Word report; requires python-docx."""
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed (pip install python-docx).") from exc
    import re

    figures = _all_figures(analysis, lang, study)
    html_text = build_html(analysis, lang, figures, img_src=lambda k: f"fig://{k}",
                           study=study)
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10)
    d.styles["Normal"].font.color.rgb = RGBColor(0x14, 0x14, 0x13)
    # the interface's look: serif headings, the title in ink, sections in terracotta
    for name, colour in (("Title", (0x14, 0x14, 0x13)), ("Heading 1", (0xC6, 0x61, 0x3F)),
                         ("Heading 2", (0x14, 0x14, 0x13))):
        style = d.styles[name]
        style.font.name = "Georgia"
        style.font.bold = False
        style.font.color.rgb = RGBColor(*colour)

    tokens = re.split(r"(<h1>.*?</h1>|<h2[^>]*>.*?</h2>|<h3>.*?</h3>|<table[^>]*>.*?</table>|"
                      r"<p[^>]*>.*?</p>|<li>.*?</li>)", html_text, flags=re.S)

    def strip(s):
        return html.unescape(re.sub(r"<[^>]+>", "", re.sub(r"<br\s*/?>", "\n", s))).strip()

    for tok in tokens:
        if tok.startswith("<h1>"):
            d.add_heading(strip(tok), level=0)
        elif tok.startswith("<h2"):
            d.add_heading(strip(tok), level=1)
        elif tok.startswith("<h3>"):
            d.add_heading(strip(tok), level=2)
        elif tok.startswith("<li>"):
            d.add_paragraph(strip(tok), style="List Bullet")
        elif tok.startswith("<table"):
            rows = re.findall(r"<tr>(.*?)</tr>", tok, flags=re.S)
            cells = [re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, flags=re.S) for r in rows]
            if cells:
                table = d.add_table(rows=len(cells), cols=len(cells[0]))
                table.style = "Light Grid Accent 1"
                for i, row in enumerate(cells):
                    for j, c in enumerate(row):
                        if j < len(table.columns):
                            table.cell(i, j).text = strip(c)
        elif tok.startswith("<p"):
            m = re.search(r"src='fig://([a-z_0-9]+)'", tok)
            if m:
                d.add_picture(io.BytesIO(figures[m.group(1)]), width=Inches(6.3))
            else:
                text = strip(tok)
                if text:
                    d.add_paragraph(text)
    d.save(path)


def export_report(path: str, analysis, lang: str, study=None) -> None:
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        export_pdf(path, analysis, lang, study)
    elif ext == "docx":
        export_docx(path, analysis, lang, study)
    else:
        export_html(path if ext in ("html", "htm") else path + ".html", analysis, lang, study)
