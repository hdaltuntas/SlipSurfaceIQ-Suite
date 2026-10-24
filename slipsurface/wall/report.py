"""
Calculation report for SlipSurface IQ Wall.

The report is assembled as HTML (the inputs, the earth pressure by every
method and the forces, the stability checks, the bearing capacity by every
method, the global stability, the earthquake of TBDY 2018, the reinforced
concrete design of every member, the bar bending schedule and the
quantities, the figures — the reinforcement drawing among them — the
warnings and the method notes) and exported as
  * PDF   — through reportlab (see `slipsurface.common.pdf`)
  * HTML  — a single self-contained file (figures embedded as base64)
  * DOCX  — optional, needs `python-docx`

There is one assembly, `build_html()`, so all three formats say the same thing.
"""

from __future__ import annotations

import base64
import datetime
import html
import io
import math
from typing import Any, Dict, List, Optional

from matplotlib.figure import Figure

from .config import APP_NAME, APP_VERSION
from .i18n import TRANSLATIONS, warning_text
from .plotting import Plotter, available_figures
from .render import figure_size
from .summary import (
    bearing_table,
    design_table,
    forces_table,
    methods_table,
    schedule_table,
    seismic_table,
    stability_table,
)

TEXTS = {
    "en": {
        "title": f"{APP_NAME} — Calculation Report",
        "date": "Date", "analyst": "Analyst", "software": "Software",
        "sec_inputs": "1. Input data", "sec_geom": "1.1 Wall geometry",
        "sec_loads": "1.2 Loads", "sec_soils": "1.3 Soils and water",
        "sec_opts": "1.4 Methods and criteria", "sec_mat": "1.5 Materials",
        "sec_earth": "2. Earth pressure and forces",
        "sec_methods": "2.1 Earth pressure on the virtual back, by method",
        "sec_forces": "2.2 Forces on the wall (per metre, moments about the toe)",
        "sec_stab": "3. Stability", "sec_checks": "3.1 Checks",
        "sec_slide": "3.2 Sliding resistance and base pressure",
        "sec_bearing": "3.3 Bearing capacity of the foundation",
        "sec_factors": "3.4 Factors of the chosen method",
        "sec_global": "3.5 Global stability (slip circles)",
        "sec_seismic": "4. Earthquake (TBDY 2018)",
        "sec_seis_par": "4.1 Seismic parameters",
        "sec_seis_forces": "4.2 Forces on the wall in the governing case",
        "sec_design": "5. Reinforced concrete design (TS 500, TBDY 2018)",
        "sec_design_mat": "5.1 Materials and combinations",
        "sec_design_sec": "5.2 Sections",
        "sec_design_cut": "5.3 Curtailment and anchorage",
        "sec_schedule": "5.4 Bar bending schedule and quantities (per metre of wall)",
        "sec_heel": "6. The heel length the stability checks need",
        "sec_figs": "7. Figures", "sec_warn": "8. Warnings", "sec_notes": "9. Method notes",
        "parameter": "Parameter", "value": "Value", "unit": "Unit",
        "H": "Height H (underside of footing to top of stem)", "t_top": "Stem thickness, top",
        "t_bot": "Stem thickness, base", "taper": "Inclined face of the stem",
        "toe": "Toe length", "heel": "Heel length", "B": "Base width B", "t_f": "Footing thickness",
        "D": "Soil over the toe", "beta": "Backslope β", "toe_slope": "Slope in front",
        "key": "Shear key (depth × width, offset from the toe)", "none": "none",
        "q_dead": "Permanent surcharge", "q_live": "Live surcharge",
        "P_top": "Vertical line load on the stem", "H_top": "Horizontal load at the top",
        "soil": "Soil", "gamma": "γ kN/m³", "gamma_sat": "γsat kN/m³", "phi": "φ' °",
        "c": "c' kPa", "hw": "Water behind the wall, above the base",
        "hf": "Water in front, above the base", "dw": "Water table below the base",
        "method": "Earth pressure", "delta": "Wall friction δ/φ", "passive": "Passive in front",
        "passive_on": "counted, {pct:.0f}% of {method}", "passive_off": "not counted",
        "base": "Base friction δb/φ · adhesion ca/c",
        "bearing_method": "Bearing capacity factors", "embedment": "Embedment as surcharge",
        "inclination": "Load inclination factors", "depth": "Depth factors",
        "q_allow": "Allowable soil pressure", "yes": "yes", "no": "no",
        "criteria": "Required factors of safety (static · seismic)",
        "grade": "Concrete", "steel": "Reinforcing steel", "cover": "Clear cover, stem · footing",
        "gc": "Unit weight of concrete",
        "resist": "Sliding resistance", "friction": "Base friction V·tan δb",
        "adhesion": "Adhesion ca·B", "pp": "Passive resistance counted",
        "q_toe": "Pressure at the toe", "q_heel": "Pressure at the heel",
        "contact": "Length in contact",
        "B_eff": "Effective width B' = B − 2e", "sigma_v": "Vertical stress σv = ΣV / B'",
        "q_ult": "Ultimate capacity q_ult", "factor": "Factor",
        "c_term": "cohesion", "q_term": "surcharge", "g_term": "self weight",
        "gl_bishop": "Bishop, simplified", "gl_fellenius": "Fellenius (same circle)",
        "gl_fellenius_min": "Fellenius, own critical circle", "gl_circle": "Critical circle",
        "gl_tried": "Circles tried",
        "mat": "fck = {fck:.0f} MPa, fcd = {fcd:.2f} MPa, fctd = {fctd:.3f} MPa, "
               "fyk = {fyk:.0f} MPa, fyd = {fyd:.1f} MPa, k1 = {k1:.3f}, ρb = {rho_b:.4f}, "
               "ρmax = {rho_max:.4f}",
        "combos": "Combinations: {names}",
        "cut_yes": "Half the bars of the back face of the stem stop {y:.2f} m above the "
                   "footing (theoretical point {t:.2f} m, extended by {ext:.2f} m).",
        "cut_no": "All the bars of the back face of the stem run the full height.",
        "anch": "Bar", "lb": "lb m", "avail": "Straight length m", "anch_state": "Anchorage",
        "straight": "straight", "hooked": "hooked end needed",
        "quant": "Concrete {v:.3f} m³/m · steel {s:.1f} kg/m · {r:.0f} kg/m³",
        "heel_found": "The shortest heel that satisfies the stability and bearing checks, static "
                      "and seismic, is {heel:.2f} m; the wall has {cur:.2f} m.",
        "heel_none": "No heel up to 3·H satisfies the stability and bearing checks.",
        "ok": "OK", "notok": "NOT OK", "na": "n/a",
        "notes": [
            "The wall is analysed per metre run. The earth pressure acts on the virtual back, "
            "the vertical plane through the end of the heel; the soil between the stem and "
            "that plane moves with the wall and counts as its weight. Rankine's thrust acts "
            "parallel to the backfill surface, Coulomb's and the trial wedge's at the wall "
            "friction δ; the pressure at rest (EN 1997-1 eq. 9.2) is for a wall that cannot "
            "yield. The trial wedge tries every plane through the heel and keeps the largest "
            "thrust; it reproduces Coulomb's coefficient and gives the critical plane.",
            "Water behind the wall pushes hydrostatically and lifts the base with a pressure "
            "that varies linearly from the water in front to the water behind; below the "
            "water table the soil counts with its buoyant unit weight for the earth pressure "
            "and with its saturated one for its weight.",
            "Sliding resists with the base friction V·tan(δb), the adhesion ca·B and, when "
            "counted, a share of the passive resistance in front (down to the bottom of the "
            "key when there is one). Overturning is taken about the toe; the uplift and the "
            "upward vertical inertia overturn. The live surcharge on the heel is left out of "
            "sliding and overturning and kept for the bearing capacity.",
            "The base pressure is a trapezoid while e ≤ B/6 and a triangle beyond. The "
            "bearing capacity is that of a strip of Meyerhof's effective width B' = B − 2e "
            "by the general equation with the factors of Terzaghi, Meyerhof, Brinch Hansen, "
            "Vesić and EN 1997-1 (SlipSurface IQ Bearing's), the load inclination, the embedment and "
            "a slope in front of the toe as chosen; it is compared with ΣV / B'.",
            "Global stability: circles passing under the footing and the key, coming out in "
            "front of the wall and in the backfill; Bishop's simplified method and "
            "Fellenius' ordinary method, on a grid of centres and tangent depths refined "
            "around the lowest factor.",
            "Members: the stem is a cantilever from the footing under the earth pressure on a "
            "vertical back of its own height (the design method's coefficient, horizontal "
            "component), the water, the surcharges and the load at its top; the toe and the "
            "heel are cantilevers from the faces of the stem under the contact pressure and "
            "the uplift against their weight and the soil and surcharge on them. The "
            "combinations are those of TS 500 (1.4G + 1.6Q + 1.6H, 0.9G + 1.6H) and TBDY "
            "2018 (G + Q + H + E, 0.9G + H + E, kv up and down).",
            "Sections one metre wide by TS 500: the tension steel from the rectangular stress "
            "block (0.85·fcd, k1), at least ρmin·b·h and at most min(0.85·ρb, 0.02)·b·d; the "
            "shear carried by the concrete alone, Vd ≤ Vcr = 0.65·fctd·b·d at d from the face "
            "of the stem (at the face for the heel); the main bars at most min(1.5·h, 200 mm) "
            "apart, the distribution bars at most 300 mm, and at least a fifth of the main "
            "steel; lb = 0.12·(fyd/fctd)·φ ≥ 20·φ, ×1.4 for top bars.",
            "Not included: settlement, the drainage of the backfill, the hydrodynamic "
            "pressure of free water in an earthquake, and the joints and the lap splices "
            "along the wall; check them separately.",
        ],
        "note_seismic": "Earthquake (TBDY 2018): Fs and F1 of the site class (Tables 2.1, "
                        "2.2), SDS = Ss·Fs, SD1 = S1·F1; the equivalent seismic coefficient of "
                        "a retaining structure kh = βr·0.4·SDS (§16.12), βr by the "
                        "displacement the wall may accept (1 when it may not), kv = ratio·kh "
                        "up and down. The dynamic increment ΔP = PAE − Pa by Mononobe–Okabe "
                        "(PAE = ½·γ·h²·(1 − kv)·KAE with the surcharge), Seed–Whitman "
                        "(ΔKAE = ¾·kh) or Wood (ΔP = kh·γ·h², unyielding walls) acts at the "
                        "chosen share of the height; the wall and the soil on its heel carry "
                        "kh·W horizontally and kv·W vertically. Stability under the earthquake "
                        "is checked against the factors of TBDY 2018 §16.7 for shallow "
                        "foundations (sliding γRh, bearing γRv) and the others given.",
    },
    "tr": {
        "title": f"{APP_NAME} — Hesap Raporu",
        "date": "Tarih", "analyst": "Hazırlayan", "software": "Yazılım",
        "sec_inputs": "1. Girdi verileri", "sec_geom": "1.1 Duvar geometrisi",
        "sec_loads": "1.2 Yükler", "sec_soils": "1.3 Zeminler ve su",
        "sec_opts": "1.4 Yöntemler ve ölçütler", "sec_mat": "1.5 Malzemeler",
        "sec_earth": "2. Toprak basıncı ve kuvvetler",
        "sec_methods": "2.1 Sanal arka yüzdeki toprak basıncı, yönteme göre",
        "sec_forces": "2.2 Duvara etkiyen kuvvetler (metre başına, momentler ön uca göre)",
        "sec_stab": "3. Duraylılık", "sec_checks": "3.1 Kontroller",
        "sec_slide": "3.2 Kayma direnci ve taban basıncı",
        "sec_bearing": "3.3 Temel zemininin taşıma gücü",
        "sec_factors": "3.4 Seçilen yöntemin katsayıları",
        "sec_global": "3.5 Genel duraylılık (kayma daireleri)",
        "sec_seismic": "4. Deprem (TBDY 2018)",
        "sec_seis_par": "4.1 Deprem parametreleri",
        "sec_seis_forces": "4.2 Belirleyici durumda duvara etkiyen kuvvetler",
        "sec_design": "5. Betonarme tasarım (TS 500, TBDY 2018)",
        "sec_design_mat": "5.1 Malzemeler ve yük birleşimleri",
        "sec_design_sec": "5.2 Kesitler",
        "sec_design_cut": "5.3 Donatı kesimi ve kenetlenme",
        "sec_schedule": "5.4 Donatı açılımı ve metrajı (duvarın metresi başına)",
        "sec_heel": "6. Duraylılık kontrollerinin gerektirdiği topuk boyu",
        "sec_figs": "7. Şekiller", "sec_warn": "8. Uyarılar", "sec_notes": "9. Yöntem notları",
        "parameter": "Parametre", "value": "Değer", "unit": "Birim",
        "H": "Yükseklik H (temel altından gövde tepesine)", "t_top": "Gövde kalınlığı, tepe",
        "t_bot": "Gövde kalınlığı, taban", "taper": "Gövdenin eğik yüzü",
        "toe": "Ön ampatman boyu", "heel": "Arka ampatman (topuk) boyu", "B": "Taban genişliği B",
        "t_f": "Temel kalınlığı", "D": "Ön ampatman üzerindeki zemin", "beta": "Arka şev β",
        "toe_slope": "Önündeki şev",
        "key": "Kayma dişi (derinlik × genişlik, ön uçtan uzaklık)", "none": "yok",
        "q_dead": "Kalıcı sürşarj", "q_live": "Hareketli sürşarj",
        "P_top": "Gövdeye düşey çizgisel yük", "H_top": "Tepede yatay yük",
        "soil": "Zemin", "gamma": "γ kN/m³", "gamma_sat": "γdoy kN/m³", "phi": "φ' °",
        "c": "c' kPa", "hw": "Duvar arkasında su, tabandan", "hf": "Önde su, tabandan",
        "dw": "Su tablasının tabandan derinliği",
        "method": "Toprak basıncı", "delta": "Duvar sürtünmesi δ/φ",
        "passive": "Öndeki pasif direnç",
        "passive_on": "hesaba katıldı, {method} değerinin %{pct:.0f}'i",
        "passive_off": "hesaba katılmadı",
        "base": "Taban sürtünmesi δb/φ · adezyon ca/c",
        "bearing_method": "Taşıma gücü katsayıları", "embedment": "Gömülme sürşarj olarak",
        "inclination": "Yük eğim katsayıları", "depth": "Derinlik katsayıları",
        "q_allow": "Zemin emniyet gerilmesi", "yes": "evet", "no": "hayır",
        "criteria": "Gerekli güvenlik sayıları (statik · deprem)",
        "grade": "Beton", "steel": "Donatı çeliği", "cover": "Net paspayı, gövde · temel",
        "gc": "Betonun birim hacim ağırlığı",
        "resist": "Kayma direnci", "friction": "Taban sürtünmesi V·tan δb",
        "adhesion": "Adezyon ca·B", "pp": "Hesaba katılan pasif direnç",
        "q_toe": "Ön uçta gerilme", "q_heel": "Arka uçta gerilme",
        "contact": "Temas eden uzunluk",
        "B_eff": "Etkin genişlik B' = B − 2e", "sigma_v": "Düşey gerilme σv = ΣV / B'",
        "q_ult": "Nihai taşıma gücü q_ult", "factor": "Katsayı",
        "c_term": "kohezyon", "q_term": "sürşarj", "g_term": "zati ağırlık",
        "gl_bishop": "Bishop, basitleştirilmiş", "gl_fellenius": "Fellenius (aynı daire)",
        "gl_fellenius_min": "Fellenius, kendi kritik dairesi", "gl_circle": "Kritik daire",
        "gl_tried": "Denenen daire",
        "mat": "fck = {fck:.0f} MPa, fcd = {fcd:.2f} MPa, fctd = {fctd:.3f} MPa, "
               "fyk = {fyk:.0f} MPa, fyd = {fyd:.1f} MPa, k1 = {k1:.3f}, ρb = {rho_b:.4f}, "
               "ρmax = {rho_max:.4f}",
        "combos": "Yük birleşimleri: {names}",
        "cut_yes": "Gövde arka yüz donatısının yarısı temel üstünden {y:.2f} m yukarıda "
                   "kesilir (kuramsal nokta {t:.2f} m, {ext:.2f} m uzatılarak).",
        "cut_no": "Gövde arka yüz donatısının tamamı tam boy devam eder.",
        "anch": "Donatı", "lb": "lb m", "avail": "Düz boy m", "anch_state": "Kenetlenme",
        "straight": "düz", "hooked": "kancalı uç gerekir",
        "quant": "Beton {v:.3f} m³/m · çelik {s:.1f} kg/m · {r:.0f} kg/m³",
        "heel_found": "Statik ve depremde duraylılık ve taşıma gücü kontrollerini sağlayan en "
                      "kısa topuk {heel:.2f} m; duvarda {cur:.2f} m.",
        "heel_none": "3·H'ye kadar hiçbir topuk boyu duraylılık ve taşıma gücü kontrollerini "
                     "sağlamıyor.",
        "ok": "UYGUN", "notok": "UYGUN DEĞİL", "na": "—",
        "notes": [
            "Duvar metre başına çözümlenir. Toprak basıncı, topuk ucundan geçen düşey düzlem "
            "olan sanal arka yüze etkir; gövde ile bu düzlem arasındaki zemin duvarla birlikte "
            "hareket eder ve ağırlığına katılır. Rankine itkisi dolgu yüzeyine paralel, "
            "Coulomb ve deneme kaması itkisi duvar sürtünme açısı δ ile etkir; sükûnetteki "
            "basınç (EN 1997-1 den. 9.2) yer değiştiremeyen duvar içindir. Deneme kaması "
            "topuktan geçen her düzlemi dener ve en büyük itkiyi alır; Coulomb katsayısını "
            "verir ve kritik düzlemi gösterir.",
            "Duvar arkasındaki su hidrostatik olarak iter ve tabanı, öndeki sudan arkadaki suya "
            "doğrusal değişen bir basınçla kaldırır; su tablası altında zemin, toprak basıncında "
            "batık, ağırlığında doygun birim hacim ağırlığıyla alınır.",
            "Kaymaya karşı direnç; taban sürtünmesi V·tan(δb), adezyon ca·B ve seçilirse "
            "öndeki pasif direncin bir payıdır (kayma dişi varsa diş altına kadar). Devrilme ön "
            "uca göre alınır; su kaldırması ve yukarı düşey atalet devirici sayılır. Topuk "
            "üzerindeki hareketli sürşarj kayma ve devrilmede dikkate alınmaz, taşıma gücünde "
            "alınır.",
            "Taban basıncı e ≤ B/6 iken yamuk, ötesinde üçgendir. Taşıma gücü, Meyerhof etkin "
            "genişliği B' = B − 2e olan şerit temel için genel denklemle, Terzaghi, Meyerhof, "
            "Brinch Hansen, Vesić ve EN 1997-1 katsayılarıyla (SlipSurface IQ Bearing'inkiler), seçilen "
            "yük eğimi, gömülme ve topuk önündeki şev ile hesaplanır ve ΣV / B' ile "
            "karşılaştırılır.",
            "Genel duraylılık: temelin ve dişin altından geçen, duvarın önünde ve arka dolguda "
            "yüzeye çıkan daireler; Bishop basitleştirilmiş yöntemi ve Fellenius yöntemi, en "
            "küçük katsayı çevresinde sıklaştırılan merkez ve teğet derinliği ızgarasıyla.",
            "Elemanlar: gövde, kendi yüksekliğindeki düşey arka yüze etkiyen toprak basıncı "
            "(tasarım yönteminin katsayısı, yatay bileşen), su, sürşarjlar ve tepedeki yük "
            "altında temelden çıkan bir konsoldur; ön ve arka ampatman, taban basıncı ve su "
            "kaldırması ile kendi ağırlıkları, üzerlerindeki zemin ve sürşarj altında gövde "
            "yüzlerinden çıkan konsollardır. Yük birleşimleri TS 500 (1.4G + 1.6Q + 1.6H, "
            "0.9G + 1.6H) ve TBDY 2018'dir (G + Q + H + E, 0.9G + H + E, kv yukarı ve aşağı).",
            "Bir metre genişlikteki kesitler TS 500'e göre: çekme donatısı dikdörtgen gerilme "
            "bloğundan (0.85·fcd, k1), en az ρmin·b·h ve en çok min(0.85·ρb, 0.02)·b·d; kesme "
            "yalnızca betonla, Vd ≤ Vcr = 0.65·fctd·b·d, gövde yüzünden d uzaklıkta (arka "
            "ampatmanda yüzde); ana donatı aralığı en çok min(1.5·h, 200 mm), dağıtma "
            "donatısı en çok 300 mm ve ana donatının en az beşte biri; lb = 0.12·(fyd/fctd)·φ ≥ "
            "20·φ, üst donatıda ×1.4.",
            "Kapsam dışı: oturma, dolgu drenajı, depremde serbest suyun hidrodinamik basıncı, "
            "duvar boyunca derzler ve bindirmeler; ayrıca kontrol ediniz.",
        ],
        "note_seismic": "Deprem (TBDY 2018): yerel zemin sınıfının Fs ve F1 katsayıları "
                        "(Tablo 2.1, 2.2), SDS = Ss·Fs, SD1 = S1·F1; istinat yapısının eşdeğer "
                        "deprem katsayısı kh = βr·0.4·SDS (§16.12), βr duvarın kabul "
                        "edilebilen yer değiştirmesine göre (edemiyorsa 1), kv = oran·kh "
                        "yukarı ve aşağı. Dinamik artış ΔP = PAE − Pa; Mononobe–Okabe "
                        "(PAE = ½·γ·h²·(1 − kv)·KAE, sürşarjla), Seed–Whitman (ΔKAE = ¾·kh) "
                        "veya Wood (ΔP = kh·γ·h², yer değiştiremeyen duvar) ile, seçilen "
                        "yükseklik payında etkir; duvar ve topuğundaki zemin yatayda kh·W, "
                        "düşeyde kv·W taşır. Depremde duraylılık, TBDY 2018 §16.7 yüzeysel "
                        "temel dayanım katsayılarına (kayma γRh, taşıma γRv) ve verilen diğer "
                        "ölçütlere göre kontrol edilir.",
    },
}

#: The interface's palette: warm paper, ink, terracotta; serif headings
_CSS = """
body { font-family: system-ui, -apple-system, 'Segoe UI', 'Helvetica Neue', Arial, sans-serif;
       font-size: 10pt; color: #141413; }
h1, h2, h3 { font-family: 'Tiempos Text', 'Source Serif 4', 'Iowan Old Style', Palatino,
             Georgia, 'DejaVu Serif', serif; font-weight: 500; }
h1 { font-size: 20pt; color: #141413; margin-bottom: 2px; }
h2 { font-size: 14pt; color: #C6613F; border-bottom: 1px solid #E3E0D5; padding-bottom: 2px;
     margin-top: 20px; }
h3 { font-size: 11.5pt; color: #141413; margin-top: 12px; }
table { border-collapse: collapse; margin: 4px 0 8px 0; }
th { background: #F0EEE6; text-align: left; padding: 3px 6px; border: 1px solid #E3E0D5;
     font-size: 9pt; }
td { padding: 3px 6px; border: 1px solid #E3E0D5; font-size: 9pt; }
.ok { color: #3F7F4F; font-weight: bold; } .bad { color: #B0413E; font-weight: bold; }
.meta { color: #73726C; } .note { color: #73726C; font-size: 9pt; }
"""


def _esc(x) -> str:
    return html.escape(str(x))


def _f(x, nd=2) -> str:
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return "∞" if x == math.inf else "—"
    return f"{x:,.{nd}f}"


def _table(headers: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None) -> str:
    out = ["<table width='100%'>"]
    if widths:
        cells = "".join(f"<th width='{w}%'>{_esc(h)}</th>" for h, w in zip(headers, widths))
    else:
        cells = "".join(f"<th>{_esc(h)}</th>" for h in headers)
    out.append("<tr>" + cells + "</tr>")
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def _kv_table(T, rows: List[List[Any]]) -> str:
    return _table([T["parameter"], T["value"], T["unit"]], rows, [50, 36, 14])


def _summary_table(tab: dict) -> str:
    """A table from `summary` with its states as coloured text."""
    rows = []
    for row in tab["rows"]:
        cells = []
        for cell, state in zip(row["cells"], row["states"]):
            text = _esc(cell)
            cells.append(f'<span class="{state}">{text}</span>' if state in ("ok", "bad")
                         else text)
        if row.get("primary"):
            cells[0] = f"<b>{cells[0]} *</b>"
        rows.append(cells)
    return _table(tab["columns"], rows)


# ----------------------------------------------------------------------
def _png(fig: Figure, dpi: int) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    return buf.getvalue()


def render_figures(wall, lang: str, dpi: int = 130) -> Dict[str, bytes]:
    """The report figures as PNG bytes (off-screen, white, untitled)."""
    plotter = Plotter(wall, lang, "paper", titles=False)
    out = {}
    for key in available_figures(wall):
        fig = Figure(figsize=figure_size(key), dpi=dpi)
        plotter.draw(key, fig)
        out[key] = _png(fig, dpi)
    return out


def build_html(wall, lang: str, figures: Dict[str, bytes], img_src=None) -> str:
    """
    Builds the report HTML. img_src(key) -> value for the <img src> attribute;
    the default embeds base64 data URIs (the self-contained HTML file). The PDF
    and DOCX writers pass a `fig://<key>` mapper and resolve the keys against
    the figure dictionary themselves.
    """
    T = TEXTS.get(lang, TEXTS["en"])
    L = TRANSLATIONS.get(lang, TRANSLATIONS["en"])
    w, res = wall, wall.results
    g = w.geo
    info = w.config.get("project_info", {})
    if img_src is None:
        def img_src(key):
            return "data:image/png;base64," + base64.b64encode(figures[key]).decode("ascii")

    def yes(flag):
        return T["yes"] if flag else T["no"]

    parts = [f"<html><head><meta charset='utf-8'><style>{_CSS}</style></head><body>"]
    parts.append(f"<h1>{_esc(info.get('title') or T['title'])}</h1>")
    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    meta = f"{T['date']}: {stamp}"
    if info.get("analyst"):
        meta += f" &nbsp;·&nbsp; {T['analyst']}: {_esc(info['analyst'])}"
    meta += f" &nbsp;·&nbsp; {T['software']}: {APP_NAME} v{APP_VERSION}"
    parts.append(f"<p class='meta'>{meta}</p>")

    # ---------------------------------------------------------------- inputs
    parts.append(f"<h2>{T['sec_inputs']}</h2>")
    parts.append(f"<h3>{T['sec_geom']}</h3>")
    key = (f"{_f(g.key_depth)} × {_f(g.key_width)}, {_f(g.key_x)}" if g.key_on else T["none"])
    parts.append(_kv_table(T, [
        [T["H"], _f(w.H), "m"], [T["t_top"], _f(w.t_top), "m"], [T["t_bot"], _f(w.t_bot), "m"],
        [T["taper"], L[f"taper_{w.taper}"], ""], [T["toe"], _f(w.toe), "m"],
        [T["heel"], _f(w.heel), "m"], [T["B"], _f(w.B), "m"], [T["t_f"], _f(w.t_f), "m"],
        [T["D"], _f(w.D), "m"], [T["beta"], _f(w.beta, 1), "°"],
        [T["toe_slope"], _f(w.toe_slope, 1), "°"], [T["key"], key, "m"]]))
    parts.append(f"<h3>{T['sec_loads']}</h3>")
    parts.append(_kv_table(T, [
        [T["q_dead"], _f(w.q_dead, 1), "kPa"], [T["q_live"], _f(w.q_live, 1), "kPa"],
        [T["P_top"], _f(w.P_top, 1), "kN/m"], [T["H_top"], _f(w.H_top, 1), "kN/m"]]))
    parts.append(f"<h3>{T['sec_soils']}</h3>")
    rows = [[L["soil_backfill"], _f(w.bf["gamma"], 1), _f(w.bf["gamma_sat"], 1),
             _f(w.bf["phi"], 1), _f(w.bf["c"], 1)],
            [L["soil_foundation"], _f(w.fd["gamma"], 1), _f(w.fd["gamma_sat"], 1),
             _f(w.fd["phi"], 1), _f(w.fd["c"], 1)]]
    parts.append(_table([T["soil"], T["gamma"], T["gamma_sat"], T["phi"], T["c"]], rows,
                        [32, 17, 17, 17, 17]))
    parts.append(_kv_table(T, [[T["hw"], _f(w.hw), "m"], [T["hf"], _f(w.hf), "m"],
                               [T["dw"], _f(w.dw), "m"]]))
    parts.append(f"<h3>{T['sec_opts']}</h3>")
    crit = w.crit
    passive = (T["passive_on"].format(pct=100.0 * w.passive_ratio,
                                     method=L[f"pp_{w.passive_method}"])
               if (w.passive_on or g.key_on) else T["passive_off"])
    rows = [[T["method"], L[f"pm_{w.method}"], ""],
            [T["delta"], _f(w.delta_ratio), ""], [T["passive"], _esc(passive), ""],
            [T["base"], f"{_f(w.base_friction)} · {_f(w.base_adhesion)}", ""],
            [T["bearing_method"], L[f"method_{w.bearing_method}"], ""],
            [T["embedment"], yes(w.use_embedment), ""],
            [T["inclination"], yes(w.use_inclination), ""],
            [T["depth"], yes(w.use_depth), ""]]
    if w.q_allow > 0:
        rows.append([T["q_allow"], _f(w.q_allow, 0), "kPa"])
    for k1, k2, label in (("FS_sliding", "FS_sliding_seis", "chk_sliding"),
                          ("FS_overturning", "FS_overturning_seis", "chk_overturning"),
                          ("FS_bearing", "FS_bearing_seis", "chk_bearing"),
                          ("FS_global", "FS_global_seis", "chk_global")):
        rows.append([f"{T['criteria']} — {L[label]}",
                     f"{_f(crit[k1])} · {_f(crit[k2])}", ""])
    rows.append([L["ecc_label"], f"B/{_f(crit['ecc_static'], 0)} · B/{_f(crit['ecc_seismic'], 0)}",
                 ""])
    parts.append(_kv_table(T, rows))
    parts.append(f"<h3>{T['sec_mat']}</h3>")
    con = w.concrete
    parts.append(_kv_table(T, [
        [T["grade"], con["grade"], ""], [T["steel"], con["steel"], ""],
        [T["cover"], f"{_f(con['cover_stem'], 0)} · {_f(con['cover_footing'], 0)}", "mm"],
        [T["gc"], _f(w.gc, 1), "kN/m³"],
        [L["rho_min_label"], _f(con["rho_min"], 4), ""],
        [L["rho_front_label"], _f(con["rho_front"], 4), ""],
        [L["rho_dist_label"], _f(con["rho_dist"], 4), ""]]))

    # ---------------------------------------------------------------- earth pressure
    parts.append(f"<h2>{T['sec_earth']}</h2>")
    parts.append(f"<h3>{T['sec_methods']}</h3>")
    parts.append(_summary_table(methods_table(w, lang)))
    parts.append(f"<h3>{T['sec_forces']}</h3>")
    parts.append(_summary_table(forces_table(w, lang)))

    # ---------------------------------------------------------------- stability
    parts.append(f"<h2>{T['sec_stab']}</h2>")
    parts.append(f"<h3>{T['sec_checks']}</h3>")
    parts.append(_summary_table(stability_table(w, lang)))
    parts.append(f"<h3>{T['sec_slide']}</h3>")
    r = res["resistance"]
    bp = res["bearing"]["pressure"]
    parts.append(_kv_table(T, [
        [T["friction"], _f(r["friction"], 1), "kN/m"], [T["adhesion"], _f(r["adhesion"], 1),
                                                        "kN/m"],
        [T["pp"], _f(r["passive"], 1), "kN/m"], [T["resist"], _f(r["R"], 1), "kN/m"],
        [T["q_toe"], _f(bp["q_toe"], 1), "kPa"], [T["q_heel"], _f(bp["q_heel"], 1), "kPa"],
        [T["contact"], _f(bp["contact"]), "m"]]))
    b = res["bearing"]
    parts.append(f"<h3>{T['sec_bearing']}</h3>")
    parts.append(_kv_table(T, [
        [T["B_eff"], _f(b["B_eff"], 3), "m"], [T["sigma_v"], _f(b.get("sigma_v"), 1), "kPa"],
        [T["q_ult"] + f" ({L['method_' + b['primary']]})", _f(b.get("q_ult"), 1), "kPa"]]))
    if b["methods"]:
        parts.append(_summary_table(bearing_table(w, lang)))
    primary = b["methods"].get(b["primary"])
    if primary:
        parts.append(f"<h3>{T['sec_factors']}</h3>")
        frows = [[L.get(f"factor_{name}", name), _f(primary["factors"][name]["c"], 3),
                  _f(primary["factors"][name]["q"], 3), _f(primary["factors"][name]["g"], 3)]
                 for name in ("depth", "inclination", "ground")]
        frows.append(["N", _f(primary["N"]["Nc"], 2), _f(primary["N"]["Nq"], 2),
                      _f(primary["N"]["Ngamma"], 2)])
        frows.append(["kPa", _f(primary["terms"]["c"], 1), _f(primary["terms"]["q"], 1),
                      _f(primary["terms"]["g"], 1)])
        parts.append(_table([T["factor"], T["c_term"], T["q_term"], T["g_term"]], frows,
                            [40, 20, 20, 20]))
    gl = res.get("global")
    if gl and gl.get("found"):
        parts.append(f"<h3>{T['sec_global']}</h3>")
        rows = []
        for tag, case in (("static", gl["static"]), ("seismic", gl.get("seismic"))):
            if not case or not case.get("found"):
                continue
            c = case["circle"]
            rows.append([L["plot_" + tag], _f(case["bishop"], 3), _f(case["fellenius"], 3),
                         _f(case["fellenius_min"], 3),
                         f"({_f(c['xc'])}, {_f(c['yc'])}), R = {_f(c['R'])}",
                         str(case["tried"])])
        parts.append(_table(["", T["gl_bishop"], T["gl_fellenius"], T["gl_fellenius_min"],
                             T["gl_circle"], T["gl_tried"]], rows, [12, 15, 15, 15, 30, 13]))

    # ---------------------------------------------------------------- seismic
    s = res["seismic"]
    if s:
        parts.append(f"<h2>{T['sec_seismic']}</h2>")
        parts.append(f"<h3>{T['sec_seis_par']}</h3>")
        parts.append(_summary_table(seismic_table(w, lang)))
        parts.append(f"<h3>{T['sec_seis_forces']}</h3>")
        parts.append(_summary_table(forces_table(w, lang, seismic=True)))

    # ---------------------------------------------------------------- design
    des = res.get("design")
    if des:
        parts.append(f"<h2>{T['sec_design']}</h2>")
        parts.append(f"<h3>{T['sec_design_mat']}</h3>")
        parts.append(f"<p>{_esc(T['mat'].format(**des['material']))}</p>")
        parts.append(f"<p>{_esc(T['combos'].format(names=', '.join(des['combinations'])))}</p>")
        parts.append(f"<h3>{T['sec_design_sec']}</h3>")
        parts.append(_summary_table(design_table(w, lang)))
        parts.append(f"<h3>{T['sec_design_cut']}</h3>")
        cut = des["curtailment"]
        text = (T["cut_yes"].format(y=cut["y_cut"] - g.t_f, t=cut["y_theory"] - g.t_f,
                                    ext=cut["extension"])
                if cut["enabled"] else T["cut_no"])
        parts.append(f"<p>{_esc(text)}</p>")
        rows = [[L[f"bar_{a['name']}"], f"Ø{a['phi']}", _f(a["lb"]), _f(a["available"]),
                 T["straight"] if a["straight"] else T["hooked"]] for a in des["anchorage"]]
        parts.append(_table([T["anch"], "Ø", T["lb"], T["avail"], T["anch_state"]], rows,
                            [34, 10, 16, 18, 22]))
        parts.append(f"<h3>{T['sec_schedule']}</h3>")
        parts.append(_summary_table(schedule_table(w, lang)))
        q = des["quantities"]
        parts.append(f"<p>{_esc(T['quant'].format(v=q['concrete'], s=q['steel'], r=q['ratio']))}"
                     "</p>")

    # ---------------------------------------------------------------- heel
    parts.append(f"<h2>{T['sec_heel']}</h2>")
    heel = res["required_heel"]
    text = (T["heel_found"].format(heel=heel, cur=w.heel) if math.isfinite(heel)
            else T["heel_none"])
    parts.append(f"<p>{_esc(text)}</p>")

    # ---------------------------------------------------------------- figures
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_figs']}</h2>")
    for fig_key in available_figures(w):
        if fig_key not in figures:
            continue
        parts.append(f"<p class='lead'>{_esc(L.get('fig_' + fig_key, fig_key))}</p>")
        parts.append(f"<p><img src='{img_src(fig_key)}' width='640'></p>")

    if res["warnings"]:
        parts.append(f"<h2>{T['sec_warn']}</h2><ul>")
        parts.extend(f"<li>{_esc(warning_text(lang, x))}</li>" for x in res["warnings"])
        parts.append("</ul>")

    parts.append(f"<h2>{T['sec_notes']}</h2><ul class='note'>")
    notes = list(T["notes"])
    if s:
        notes.append(T["note_seismic"])
    parts.extend(f"<li>{_esc(n)}</li>" for n in notes)
    parts.append("</ul>")
    parts.append("</body></html>")
    return "\n".join(parts)


# ----------------------------------------------------------------------
# exporters
# ----------------------------------------------------------------------
def export_html(path: str, wall, lang: str) -> None:
    figures = render_figures(wall, lang)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(build_html(wall, lang, figures))


def export_pdf(path: str, wall, lang: str) -> None:
    """PDF through reportlab; the figures travel as `fig://<key>` references."""
    from ..common import pdf as pdf_writer

    T = TEXTS.get(lang, TEXTS["en"])
    info = wall.config.get("project_info", {})
    figures = render_figures(wall, lang)
    html_text = build_html(wall, lang, figures, img_src=lambda k: f"fig://{k}")
    pdf_writer.html_to_pdf(path, html_text, figures,
                           title=info.get("title", T["title"]),
                           author=info.get("analyst", ""),
                           footer=f"{info.get('title', '')} — {APP_NAME} v{APP_VERSION}")


def export_docx(path: str, wall, lang: str) -> None:
    """Word report; requires python-docx."""
    try:
        import docx
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed (pip install python-docx).") from exc
    import re

    figures = render_figures(wall, lang)
    html_text = build_html(wall, lang, figures, img_src=lambda k: f"fig://{k}")
    d = docx.Document()
    from docx.shared import RGBColor
    d.styles["Normal"].font.size = Pt(10)
    d.styles["Normal"].font.color.rgb = RGBColor(0x14, 0x14, 0x13)
    # the interface's look: serif headings, the title in ink, sections in terracotta
    for name, colour in (("Title", (0x14, 0x14, 0x13)), ("Heading 1", (0xC6, 0x61, 0x3F)),
                         ("Heading 2", (0x14, 0x14, 0x13))):
        style = d.styles[name]
        style.font.name = "Georgia"
        style.font.bold = False
        style.font.color.rgb = RGBColor(*colour)

    tokens = re.split(r"(<h1>.*?</h1>|<h2[^>]*>.*?</h2>|<h3>.*?</h3>|<table[^>]*>.*?</table>|"
                      r"<p[^>]*>.*?</p>|<li>.*?</li>)", html_text, flags=re.S)

    def strip(s):
        return html.unescape(re.sub(r"<[^>]+>", "", re.sub(r"<br\s*/?>", "\n", s))).strip()

    for tok in tokens:
        if tok.startswith("<h1>"):
            d.add_heading(strip(tok), level=0)
        elif tok.startswith("<h2"):
            d.add_heading(strip(tok), level=1)
        elif tok.startswith("<h3>"):
            d.add_heading(strip(tok), level=2)
        elif tok.startswith("<li>"):
            d.add_paragraph(strip(tok), style="List Bullet")
        elif tok.startswith("<table"):
            rows = re.findall(r"<tr>(.*?)</tr>", tok, flags=re.S)
            cells = [re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, flags=re.S) for r in rows]
            if cells:
                table = d.add_table(rows=len(cells), cols=max(len(c) for c in cells))
                table.style = "Light Grid Accent 1"
                for i, row in enumerate(cells):
                    for j, c in enumerate(row):
                        if j < len(table.columns):
                            table.cell(i, j).text = strip(c)
        elif tok.startswith("<p"):
            m = re.search(r"src='fig://([a-z_]+)'", tok)
            if m:
                d.add_picture(io.BytesIO(figures[m.group(1)]), width=Inches(6.3))
            else:
                text = strip(tok)
                if text:
                    d.add_paragraph(text)
    d.save(path)


def export_report(path: str, wall, lang: str) -> None:
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        export_pdf(path, wall, lang)
    elif ext == "docx":
        export_docx(path, wall, lang)
    else:
        export_html(path if ext in ("html", "htm") else path + ".html", wall, lang)
