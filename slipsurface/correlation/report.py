"""
Calculation report for SlipSurface IQ Correlation.

The report is assembled as HTML (inputs, index properties, the selected
parameters, every correlation layer by layer, the processed field tests,
figures, warnings and the references) and exported as
  * PDF   — through reportlab (see `slipsurface.correlation.pdf`)
  * HTML  — a single self-contained file (figures embedded as base64)
  * DOCX  — optional, needs `python-docx`

There is one assembly, `build_html()`, so all three formats say the same thing.
"""

from __future__ import annotations

import base64
import datetime
import html
import io
from typing import Any, Dict, List, Optional

from matplotlib.figure import Figure

from . import summary
from .config import APP_NAME, APP_VERSION, PARAMS
from .correlations import METHODS, reference
from .i18n import TRANSLATIONS, warning_text
from .plotting import Plotter, available

TEXTS = {
    "en": {
        "title": f"{APP_NAME} — Soil Parameter Report",
        "date": "Date", "analyst": "Analyst", "software": "Software",
        "sec_inputs": "1. Input data", "sec_site": "1.1 Site and test settings",
        "sec_soil": "1.2 Soil profile and laboratory results",
        "sec_tests": "1.3 Field tests",
        "sec_index": "2. Index properties and classification",
        "sec_selected": "3. Selected parameters",
        "sec_by_layer": "4. The correlations, layer by layer",
        "sec_processed": "5. Field tests as processed",
        "sec_figs": "6. Figures", "sec_warn": "7. Warnings", "sec_notes": "8. Method notes",
        "sec_refs": "9. Correlations used",
        "parameter": "Parameter", "value": "Value", "unit": "Unit",
        "zw": "Water table depth", "gw": "Unit weight of water γw",
        "er": "SPT hammer energy ratio ER", "borehole": "Borehole diameter",
        "sampler": "Sampler", "stickup": "Rod above the ground", "cn": "Overburden correction",
        "cn_max": "Upper limit of CN", "a": "Net area ratio of the cone a",
        "nkt": "Cone factor Nkt", "ndu": "Pore pressure factor NΔu",
        "combine": "Combination of the correlations", "eu_ratio": "Eu / cu",
        "chosen": "Correlations chosen by the user", "none": "none",
        "test": "Test", "count": "Count", "range": "Depth range",
        "layer": "Layer", "type": "Type", "top": "Top", "bottom": "Bottom",
        "zone": "SBT zone (CPT)", "mean_ic": "Mean Ic",
        "every": "Every {n}th reading is listed; the figures and the averages use all {total}.",
        "notes": [
            "Stresses: σv0 from γ above the water table and γsat below it, u0 hydrostatic from "
            "the water table, σ'v0 = σv0 − u0. A unit weight that was not entered is taken from "
            "the phase relations (w, Gs, e0) and the CPT (Robertson & Cabal 2010), combined; "
            "failing both, a typical value.",
            "SPT: N60 = N·(ER/60)·CB·CS·CR with CB from the borehole diameter (Skempton 1986), "
            "CS = 1.2 for a sampler without liners and CR from the rod length (Youd et al. "
            "2001); (N1)60 = CN·N60.",
            "CPT: qt = qc + u2·(1 − a); Qtn, Fr, Bq and Ic after Robertson (2009), with the "
            "stress exponent n found by iteration; the soil behaviour type zone from Ic.",
            "Every correlation is evaluated at each test depth and averaged over the layer "
            "(geometrically for k and cv); laboratory correlations give one value per layer.",
            "Derived parameters, in order: OCR = σ'p/σ'v0 at mid-layer; cu = 0.22·σ'p (Mesri "
            "1975) added to the undrained strengths; K0 = (1 − sin φ')·OCR^sin φ'; Eu = "
            "(Eu/cu)·cu; Cr = 0.15·Cc and Cα = 0.04·Cc added to their correlations; "
            "G0 = (γ/g)·Vs².",
            "Selection: the correlation chosen for a parameter, or all that apply, combined by "
            "the median, the mean, or a cautious estimate of the mean — Xm − 0.5·s for a "
            "strength or a stiffness, Xm + 0.5·s for a compressibility (Schneider 1997), kept "
            "within the range of the correlations.",
            "Correlations are empirical: each was fitted to a particular set of soils and "
            "carries a scatter of its own. The values here are a first estimate to be checked "
            "against the site's own tests and judgement, not a substitute for them.",
        ],
    },
    "tr": {
        "title": f"{APP_NAME} — Zemin Parametreleri Raporu",
        "date": "Tarih", "analyst": "Hazırlayan", "software": "Yazılım",
        "sec_inputs": "1. Girdi verileri", "sec_site": "1.1 Saha ve deney ayarları",
        "sec_soil": "1.2 Zemin profili ve laboratuvar sonuçları",
        "sec_tests": "1.3 Arazi deneyleri",
        "sec_index": "2. İndeks özellikleri ve sınıflandırma",
        "sec_selected": "3. Seçilen parametreler",
        "sec_by_layer": "4. Tabaka tabaka korelasyonlar",
        "sec_processed": "5. İşlenmiş arazi deneyleri",
        "sec_figs": "6. Şekiller", "sec_warn": "7. Uyarılar", "sec_notes": "8. Yöntem notları",
        "sec_refs": "9. Kullanılan korelasyonlar",
        "parameter": "Parametre", "value": "Değer", "unit": "Birim",
        "zw": "Su tablası derinliği", "gw": "Suyun birim hacim ağırlığı γw",
        "er": "SPT tokmak enerji oranı ER", "borehole": "Sondaj çapı",
        "sampler": "Numune alıcı", "stickup": "Zemin üstündeki tij boyu",
        "cn": "Örtü yükü düzeltmesi", "cn_max": "CN üst sınırı",
        "a": "Koninin net alan oranı a", "nkt": "Koni faktörü Nkt",
        "ndu": "Boşluk suyu faktörü NΔu",
        "combine": "Korelasyonların birleştirilmesi", "eu_ratio": "Eu / cu",
        "chosen": "Kullanıcının seçtiği korelasyonlar", "none": "yok",
        "test": "Deney", "count": "Adet", "range": "Derinlik aralığı",
        "layer": "Tabaka", "type": "Tür", "top": "Üst", "bottom": "Alt",
        "zone": "ZDT bölgesi (CPT)", "mean_ic": "Ortalama Ic",
        "every": "Her {n}. okuma listelenmiştir; şekiller ve ortalamalar tüm {total} okumayı "
                 "kullanır.",
        "notes": [
            "Gerilmeler: σv0 su tablası üstünde γ, altında γdoy ile; u0 su tablasından "
            "hidrostatik; σ'v0 = σv0 − u0. Girilmeyen birim hacim ağırlığı faz bağıntılarından "
            "(w, Gs, e0) ve CPT'den (Robertson & Cabal 2010) birleştirilerek alınır; ikisi de "
            "yoksa tipik bir değer kullanılır.",
            "SPT: N60 = N·(ER/60)·CB·CS·CR; CB sondaj çapından (Skempton 1986), gömleksiz numune "
            "alıcıda CS = 1.2, CR tij boyundan (Youd vd. 2001); (N1)60 = CN·N60.",
            "CPT: qt = qc + u2·(1 − a); Qtn, Fr, Bq ve Ic Robertson'a (2009) göre, gerilme üssü "
            "n iterasyonla; zemin davranış türü bölgesi Ic'den.",
            "Her korelasyon her deney derinliğinde hesaplanır ve tabaka boyunca ortalanır (k ve "
            "cv için geometrik); laboratuvar korelasyonları tabaka başına tek değer verir.",
            "Türetilen parametreler, sırasıyla: tabaka ortasında OCR = σ'p/σ'v0; drenajsız "
            "dayanımlara cu = 0.22·σ'p (Mesri 1975) eklenir; K0 = (1 − sin φ')·OCR^sin φ'; "
            "Eu = (Eu/cu)·cu; Cr = 0.15·Cc ve Cα = 0.04·Cc kendi korelasyonlarına eklenir; "
            "G0 = (γ/g)·Vs².",
            "Seçim: parametre için seçilen korelasyon ya da uygulanabilen tümünün medyanı, "
            "ortalaması veya ihtiyatlı ortalama tahmini — dayanım ve rijitlikte Xm − 0.5·s, "
            "sıkışabilirlikte Xm + 0.5·s (Schneider 1997), korelasyonların aralığında kalacak "
            "şekilde.",
            "Korelasyonlar ampiriktir: her biri belirli zeminlere uydurulmuştur ve kendi "
            "saçılmasını taşır. Buradaki değerler, sahanın kendi deneyleri ve mühendislik "
            "muhakemesiyle denetlenecek bir ilk tahmindir; onların yerini tutmaz.",
        ],
    },
}

#: The interface's palette: warm paper, ink, terracotta; serif headings
_CSS = """
body { font-family: system-ui, -apple-system, 'Segoe UI', 'Helvetica Neue', Arial, sans-serif;
       font-size: 10pt; color: #141413; }
h1, h2, h3 { font-family: 'Tiempos Text', 'Source Serif 4', 'Iowan Old Style', Palatino,
             Georgia, 'DejaVu Serif', serif; font-weight: 500; }
h1 { font-size: 20pt; color: #141413; margin-bottom: 2px; }
h2 { font-size: 14pt; color: #C6613F; border-bottom: 1px solid #E3E0D5; padding-bottom: 2px;
     margin-top: 20px; }
h3 { font-size: 11.5pt; color: #141413; margin-top: 12px; }
table { border-collapse: collapse; margin: 4px 0 8px 0; }
th { background: #F0EEE6; text-align: left; padding: 3px 6px; border: 1px solid #E3E0D5;
     font-size: 9pt; }
td { padding: 3px 6px; border: 1px solid #E3E0D5; font-size: 9pt; }
.ok { color: #3F7F4F; font-weight: bold; } .bad { color: #B0413E; font-weight: bold; }
.meta { color: #73726C; } .note { color: #73726C; font-size: 9pt; }
"""

#: At most this many CPT readings are listed in the report
CPT_ROWS = 80


def _esc(x) -> str:
    return html.escape(str(x))


def _f(x, nd=2) -> str:
    if x is None:
        return "—"
    return f"{x:,.{nd}f}"


def _m(x, nd=2) -> str:
    """A laboratory value; zero was "not measured"."""
    return _f(float(x), nd) if x else "—"


def _table(headers: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None) -> str:
    out = ["<table width='100%'>"]
    if widths:
        cells = "".join(f"<th width='{w}%'>{_esc(h)}</th>" for h, w in zip(headers, widths))
    else:
        cells = "".join(f"<th>{_esc(h)}</th>" for h in headers)
    out.append("<tr>" + cells + "</tr>")
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def _kv_table(T, rows: List[List[Any]]) -> str:
    return _table([T["parameter"], T["value"], T["unit"]], rows, [50, 36, 14])


def _esc_table(headers, rows, widths=None) -> str:
    return _table(headers, [[_esc(c) for c in row] for row in rows], widths)


# ----------------------------------------------------------------------
def _png(fig: Figure, dpi: int) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    return buf.getvalue()


def render_figures(analysis, lang: str, keys=None, dpi: int = 130) -> Dict[str, bytes]:
    """The report figures as PNG bytes (off-screen, white, untitled)."""
    plotter = Plotter(analysis, lang, "paper", titles=False)
    out = {}
    for key in keys if keys is not None else available(analysis):
        fig = Figure(figsize=(10, 7), dpi=dpi)
        plotter.draw(key, fig)
        out[key] = _png(fig, dpi)
    return out


def _split(params: List[str]) -> List[List[str]]:
    """The selected parameters in two tables, so each fits the width of a page."""
    half = (len(params) + 1) // 2
    return [params[:half], params[half:]] if len(params) > 9 else [params]


def build_html(analysis, lang: str, figures: Dict[str, bytes], img_src=None) -> str:
    """
    Builds the report HTML. img_src(key) -> value for the <img src> attribute;
    the default embeds base64 data URIs (the self-contained HTML file). The PDF
    and DOCX writers pass a `fig://<key>` mapper and resolve the keys against
    the figure dictionary themselves.
    """
    T = TEXTS.get(lang, TEXTS["en"])
    L = TRANSLATIONS.get(lang, TRANSLATIONS["en"])
    a, res = analysis, analysis.results
    layers = res["layers"]
    info = a.config.get("project_info", {})
    if img_src is None:
        def img_src(key):
            return "data:image/png;base64," + base64.b64encode(figures[key]).decode()

    parts = [f"<html><head><meta charset='utf-8'><style>{_CSS}</style></head><body>"]
    parts.append(f"<h1>{_esc(info.get('title', ''))}</h1>")
    parts.append(f"<p class='meta'>{_esc(T['title'])}<br>"
                 f"{T['date']}: {datetime.date.today().isoformat()} &nbsp;|&nbsp; "
                 f"{T['analyst']}: {_esc(info.get('analyst', ''))} &nbsp;|&nbsp; "
                 f"{T['software']}: {APP_NAME} v{APP_VERSION}</p>")

    # ---------------- 1. inputs
    parts.append(f"<h2>{T['sec_inputs']}</h2>")
    parts.append(f"<h3>{T['sec_site']}</h3>")
    chosen = "; ".join(f"{L['sym_' + p]}: {reference(k, lang)}" for p, k in a.chosen.items())
    parts.append(_kv_table(T, [
        [T["zw"], _f(a.zw), "m"], [T["gw"], _f(a.gw), "kN/m³"],
        [T["er"], _f(a.spt["energy_ratio"], 0), "%"], [T["borehole"], _f(a.spt["borehole"], 0), "mm"],
        [T["sampler"], L[f"sampler_{a.spt['sampler']}"], ""],
        [T["stickup"], _f(a.spt["stickup"]), "m"],
        [T["cn"], L[f"cn_{a.spt['cn_method']}"], ""], [T["cn_max"], _f(a.spt["cn_max"]), ""],
        [T["a"], _f(a.cpt["area_ratio"]), ""], [T["nkt"], _f(a.cpt["Nkt"], 1), ""],
        [T["ndu"], _f(a.cpt["Ndu"], 1), ""],
        [T["combine"], L[f"combine_{a.combine}"], ""], [T["eu_ratio"], _f(a.Eu_ratio, 0), ""],
        [T["chosen"], _esc(chosen) if chosen else T["none"], ""],
    ]))

    parts.append(f"<h3>{T['sec_soil']}</h3>")
    rows = []
    for i, lay in enumerate(a.layers, 1):
        rows.append([i, _esc(lay["name"]), L["behaviour_" + lay["behaviour"]], _f(lay["top"]),
                     _f(lay["bottom"])] + [_m(lay.get(k), 1)
                                           for k in ("gravel", "fines", "LL", "PL", "w")]
                    + [_m(lay.get("Gs")), _m(lay.get("D10"), 3), _m(lay.get("qu"), 0)])
    parts.append(_table(["#", T["layer"], T["type"], f"{T['top']} (m)", f"{T['bottom']} (m)",
                         L["col_gravel"], L["col_fines"], L["col_LL"], L["col_PL"], L["col_w"],
                         "Gs", L["col_D10"], L["col_qu"]], rows,
                        [3, 17, 10, 6, 6, 7, 7, 6, 6, 6, 6, 9, 9]))

    parts.append(f"<h3>{T['sec_tests']}</h3>")
    rows = []
    for kind, label in (("spt", "SPT"), ("cpt", "CPT(u)"), ("vs", "Vs")):
        tests = res[kind]
        if tests:
            rows.append([label, len(tests), f"{tests[0]['depth']:.2f} – {tests[-1]['depth']:.2f} m"])
    if rows:
        parts.append(_table([T["test"], T["count"], T["range"]], rows, [30, 20, 50]))

    # ---------------- 2. index
    parts.append(f"<h2>{T['sec_index']}</h2>")
    rows = []
    for lay in layers:
        rows.append([_esc(lay["name"]), _esc(lay["uscs"]), _f(lay["PI"], 0), _f(lay["LI"]),
                     _f(lay["IC"]), summary.fmt("e0", lay["params"]["e0"]["value"]),
                     _f(lay["Ic"]), _esc(summary.zone_text(L, lay["zone"]))])
    parts.append(_table([T["layer"], "USCS", "PI", "LI", "IC", "e0", T["mean_ic"], T["zone"]],
                        rows, [22, 14, 7, 7, 7, 8, 9, 26]))

    # ---------------- 3. selected parameters
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_selected']}</h2>")
    shown = summary.shown_params(analysis)
    for block in _split(shown):
        headers = [T["layer"]] + [f"{L['sym_' + p]} ({summary.unit(p, L)})"
                                  if summary.unit(p, L) else L["sym_" + p] for p in block]
        rows = [[_esc(lay["name"])] + [summary.fmt(p, lay["params"][p]["value"]) for p in block]
                for lay in layers]
        parts.append(_table(headers, rows))

    # ---------------- 4. by layer
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_by_layer']}</h2>")
    for table in summary.method_tables(analysis, lang):
        parts.append(f"<h3>{_esc(table['title'])}</h3>")
        body = [[f"<b>{_esc(r[0])}</b>" if r[0] else "", f"<b>{_esc(r[1])}</b>" if r[0]
                 else _esc(r[1]), _esc(r[2]), _esc(r[3]), _esc(r[4]), _esc(r[5])]
                for r in table["rows"]]
        parts.append(_table(table["columns"], body, [11, 13, 25, 11, 5, 35]))

    # ---------------- 5. processed tests
    tables = summary.test_tables(analysis, lang)
    if tables:
        parts.append(f"<h2 style='page-break-before:always'>{T['sec_processed']}</h2>")
        for table in tables:
            parts.append(f"<h3>{_esc(table['title'])}</h3>")
            rows = table["rows"]
            if len(rows) > CPT_ROWS:
                step = -(-len(rows) // CPT_ROWS)
                parts.append(f"<p class='note'>{_esc(T['every'].format(n=step, total=len(rows)))}</p>")
                rows = rows[::step]
            parts.append(_esc_table(table["columns"], rows))

    # ---------------- 6. figures
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_figs']}</h2>")
    for key in figures:
        parts.append(f"<p class='lead'><b>{_esc(L.get(f'fig_{key}', key))}</b></p>")
        parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")

    # ---------------- 7/8/9. warnings, notes, references
    if res.get("warnings"):
        parts.append(f"<h2>{T['sec_warn']}</h2><ul>")
        parts.extend(f"<li>{_esc(warning_text(lang, w))}</li>" for w in res["warnings"])
        parts.append("</ul>")
    parts.append(f"<h2>{T['sec_notes']}</h2><ul class='note'>")
    parts.extend(f"<li>{_esc(n)}</li>" for n in T["notes"])
    parts.append("</ul>")

    used = {key for lay in layers for p in PARAMS for key in lay["params"][p]["methods"]}
    rows = [[_esc(L["sym_" + m.param]), _esc(reference(m.key, lang)),
             _esc(L[f"source_{m.source}"]), _esc(m.formula)]
            for m in METHODS.values() if m.key in used]
    parts.append(f"<h2>{T['sec_refs']}</h2>")
    parts.append(_table([L["head_param"], L["head_ref"], L["head_source"], L["head_formula"]],
                        rows, [9, 30, 13, 48]))
    parts.append("</body></html>")
    return "\n".join(parts)


# ----------------------------------------------------------------------
# exporters
# ----------------------------------------------------------------------
def export_html(path: str, analysis, lang: str) -> None:
    figures = render_figures(analysis, lang)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(build_html(analysis, lang, figures))


def export_pdf(path: str, analysis, lang: str) -> None:
    """PDF through reportlab; the figures travel as `fig://<key>` references."""
    from . import pdf as pdf_writer

    T = TEXTS.get(lang, TEXTS["en"])
    info = analysis.config.get("project_info", {})
    figures = render_figures(analysis, lang)
    html_text = build_html(analysis, lang, figures, img_src=lambda k: f"fig://{k}")
    pdf_writer.html_to_pdf(path, html_text, figures,
                           title=info.get("title", T["title"]),
                           author=info.get("analyst", ""),
                           footer=f"{info.get('title', '')} — {APP_NAME} v{APP_VERSION}")


def export_docx(path: str, analysis, lang: str) -> None:
    """Word report; requires python-docx."""
    try:
        import docx
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed (pip install python-docx).") from exc
    import re

    figures = render_figures(analysis, lang)
    html_text = build_html(analysis, lang, figures, img_src=lambda k: f"fig://{k}")
    d = docx.Document()
    from docx.shared import RGBColor
    d.styles["Normal"].font.size = Pt(10)
    d.styles["Normal"].font.color.rgb = RGBColor(0x14, 0x14, 0x13)
    # the interface's look: serif headings, the title in ink, sections in terracotta
    for name, colour in (("Title", (0x14, 0x14, 0x13)), ("Heading 1", (0xC6, 0x61, 0x3F)),
                         ("Heading 2", (0x14, 0x14, 0x13))):
        style = d.styles[name]
        style.font.name = "Georgia"
        style.font.bold = False
        style.font.color.rgb = RGBColor(*colour)

    # a small HTML -> docx walker (headings, paragraphs, tables, lists, images)
    tokens = re.split(r"(<h1>.*?</h1>|<h2[^>]*>.*?</h2>|<h3>.*?</h3>|<table[^>]*>.*?</table>|"
                      r"<p[^>]*>.*?</p>|<li>.*?</li>)", html_text, flags=re.S)

    def strip(s):
        return html.unescape(re.sub(r"<[^>]+>", "", re.sub(r"<br\s*/?>", "\n", s))).strip()

    for tok in tokens:
        if tok.startswith("<h1>"):
            d.add_heading(strip(tok), level=0)
        elif tok.startswith("<h2"):
            d.add_heading(strip(tok), level=1)
        elif tok.startswith("<h3>"):
            d.add_heading(strip(tok), level=2)
        elif tok.startswith("<li>"):
            d.add_paragraph(strip(tok), style="List Bullet")
        elif tok.startswith("<table"):
            rows = re.findall(r"<tr>(.*?)</tr>", tok, flags=re.S)
            cells = [re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, flags=re.S) for r in rows]
            if cells:
                table = d.add_table(rows=len(cells), cols=len(cells[0]))
                table.style = "Light Grid Accent 1"
                for i, row in enumerate(cells):
                    for j, c in enumerate(row):
                        if j < len(table.columns):
                            table.cell(i, j).text = strip(c)
        elif tok.startswith("<p"):
            m = re.search(r"src='fig://([a-z_]+)'", tok)
            if m:
                d.add_picture(io.BytesIO(figures[m.group(1)]), width=Inches(6.3))
            else:
                text = strip(tok)
                if text:
                    d.add_paragraph(text)
    d.save(path)


def export_report(path: str, analysis, lang: str) -> None:
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        export_pdf(path, analysis, lang)
    elif ext == "docx":
        export_docx(path, analysis, lang)
    else:
        export_html(path if ext in ("html", "htm") else path + ".html", analysis, lang)
